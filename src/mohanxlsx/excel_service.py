from __future__ import annotations

from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
import pandas as pd

from .project_config import ColumnRule


class ExcelServiceError(RuntimeError):
    pass


def read_table(path: str | Path, sheet_name: str = "") -> pd.DataFrame:
    source = Path(path)
    suffix = source.suffix.lower()

    try:
        if suffix == ".csv":
            return pd.read_csv(source, header=None, dtype=object)
        if suffix in {".xls", ".xlsx", ".xlsm"}:
            selected_sheet = sheet_name if sheet_name else 0
            return pd.read_excel(source, header=None, sheet_name=selected_sheet, dtype=object)
    except Exception as exc:
        raise ExcelServiceError(f"读取文件失败: {source.name}，{exc}") from exc

    raise ExcelServiceError(f"不支持的文件类型: {source.suffix}")


def build_dataframe(raw_df: pd.DataFrame, header_row: int) -> pd.DataFrame:
    if raw_df.empty:
        return pd.DataFrame()

    header_index = max(0, min(header_row - 1, len(raw_df.index) - 1))
    raw_headers = raw_df.iloc[header_index].tolist()
    columns = deduplicate_headers(raw_headers)

    dataframe = raw_df.iloc[header_index + 1 :].reset_index(drop=True).copy()
    dataframe.columns = columns
    return dataframe


def load_prepared_dataframe(path: str | Path, header_row: int, sheet_name: str = "") -> pd.DataFrame:
    raw_df = read_table(path, sheet_name=sheet_name)
    return build_dataframe(raw_df, header_row=header_row)


def clean_dataframe(
    dataframe: pd.DataFrame,
    remove_empty_rows: bool = True,
    drop_duplicates: bool = False,
) -> pd.DataFrame:
    cleaned = dataframe.copy()
    if remove_empty_rows:
        cleaned = cleaned.dropna(how="all")
    if drop_duplicates:
        cleaned = cleaned.drop_duplicates()
    return cleaned.reset_index(drop=True)


def apply_column_rules(dataframe: pd.DataFrame, column_rules: list[ColumnRule]) -> pd.DataFrame:
    if dataframe.empty:
        return dataframe.copy()

    enabled_columns = [rule.original_name for rule in column_rules if rule.enabled and rule.original_name in dataframe.columns]
    if enabled_columns:
        transformed = dataframe.loc[:, enabled_columns].copy()
    else:
        transformed = dataframe.copy()

    rename_map = {}
    for rule in column_rules:
        if rule.original_name in transformed.columns and rule.display_name.strip():
            rename_map[rule.original_name] = rule.display_name.strip()
    return transformed.rename(columns=rename_map)


def export_dataframe(dataframe: pd.DataFrame, path: str | Path) -> None:
    target = Path(path)
    suffix = target.suffix.lower()

    try:
        if suffix == ".csv":
            dataframe.to_csv(target, index=False, encoding="utf-8-sig")
            return
        if suffix in {".xlsx", ".xlsm"}:
            dataframe.to_excel(target, index=False)
            _format_exported_excel(target)
            return
        if suffix == ".xls":
            raise ExcelServiceError("当前仅支持导出为 .xlsx、.xlsm 或 .csv")
    except ExcelServiceError:
        raise
    except Exception as exc:
        raise ExcelServiceError(f"保存文件失败: {target.name}，{exc}") from exc

    raise ExcelServiceError(f"不支持的导出类型: {target.suffix}")


def split_and_export(
    dataframe: pd.DataFrame,
    split_columns: list[str],
    output_dir: str | Path,
    output_prefix: str,
    export_format: str,
    column_rules: list[ColumnRule],
    values: list[str] | None = None,
) -> list[Path]:
    if not split_columns:
        raise ExcelServiceError("拆分列不能为空")
    for col in split_columns:
        if col not in dataframe.columns:
            raise ExcelServiceError(f"拆分列不存在：{col}，请先重新加载或检查标题行设置")

    selected_values = {value.strip() for value in (values or []) if value.strip()}
    exported: list[Path] = []
    output_root = Path(output_dir)
    output_root.mkdir(parents=True, exist_ok=True)

    filled = dataframe.copy()
    for col in split_columns:
        filled[col] = filled[col].fillna("空值")

    for group_value, group_df in filled.groupby(split_columns, dropna=False, sort=True):
        if isinstance(group_value, tuple):
            display_value = "_".join(str(v).strip() or "空值" for v in group_value)
        else:
            display_value = str(group_value).strip() or "空值"
        if selected_values and display_value not in selected_values:
            continue
        transformed = apply_column_rules(group_df.reset_index(drop=True), column_rules)
        safe_name = _safe_filename(display_value)
        output_path = output_root / f"{output_prefix}_{safe_name}.{_normalize_export_format(export_format)}"
        export_dataframe(transformed, output_path)
        exported.append(output_path)

    return exported


def merge_files(
    source_files: list[str],
    header_row: int,
    sheet_name: str = "",
    remove_empty_rows: bool = True,
    drop_duplicates: bool = False,
) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for file_path in source_files:
        frame = load_prepared_dataframe(file_path, header_row=header_row, sheet_name=sheet_name)
        frame = clean_dataframe(
            frame,
            remove_empty_rows=remove_empty_rows,
            drop_duplicates=False,
        )
        frames.append(frame)

    if not frames:
        return pd.DataFrame()

    merged = pd.concat(frames, ignore_index=True, sort=False)
    return clean_dataframe(
        merged,
        remove_empty_rows=remove_empty_rows,
        drop_duplicates=drop_duplicates,
    )


def deduplicate_headers(headers: list[object]) -> list[str]:
    counts: dict[str, int] = {}
    result: list[str] = []

    for index, value in enumerate(headers, start=1):
        text = "" if pd.isna(value) else str(value).strip()  # type: ignore[arg-type]
        base = text or f"未命名列{index}"
        count = counts.get(base, 0)
        counts[base] = count + 1
        result.append(base if count == 0 else f"{base}_{count + 1}")

    return result


def _safe_filename(value: str) -> str:
    safe = re.sub(r"[\\/:*?\"<>|]+", "_", value).strip()
    return safe[:80] or "空值"


def _normalize_export_format(export_format: str) -> str:
    normalized = export_format.lower().lstrip(".")
    if normalized in {"xlsx", "xlsm", "csv"}:
        return normalized
    return "xlsx"


def _format_exported_excel(path: Path) -> None:
    from openpyxl.cell.cell import Cell

    workbook = load_workbook(path)
    worksheet = workbook.active
    if worksheet is None:
        return

    border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    max_width = 35
    min_width = 10

    for row_index, row in enumerate(worksheet.iter_rows(), start=1):
        for cell in row:
            cell.border = border
            if row_index == 1:
                cell.font = Font(bold=True)
                cell.alignment = header_alignment
            else:
                cell.alignment = body_alignment

    for column_cells in worksheet.columns:
        first_cell = column_cells[0]
        if not isinstance(first_cell, Cell):
            continue
        column_letter = first_cell.column_letter
        width = _measure_column_width(column_cells)
        worksheet.column_dimensions[column_letter].width = min(max(width, min_width), max_width)

    for row in worksheet.iter_rows():
        requires_wrap = any(
            cell.value is not None and len(str(cell.value)) > 30
            for cell in row
        )
        if requires_wrap:
            first_cell = row[0]
            if not isinstance(first_cell, Cell):
                continue
            worksheet.row_dimensions[first_cell.row].height = 36

    workbook.save(path)


def _measure_column_width(column_cells) -> int:
    max_length = 0
    for cell in column_cells:
        if cell.value is None:
            continue
        text = str(cell.value)
        line_length = max((len(part) for part in text.splitlines()), default=0)
        max_length = max(max_length, line_length)
    return max_length + 2
